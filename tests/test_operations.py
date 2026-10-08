import os
import unittest
from datetime import datetime, time, timedelta
from io import BytesIO
from unittest.mock import patch

os.environ['DATABASE_URL']='sqlite:///:memory:'
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from database.db import Base, SessionLocal
from database.models import User, Client, Activity, Visit, PromiseToPay, CollectionActivity, VerificationOTP, AuditLog
from main import app
from modules.security import hash_password
from modules.operations import business_today


class OperationsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
        SessionLocal.configure(bind=cls.engine)
        cls.password_hash=hash_password('Test-only-password-123')

    def setUp(self):
        Base.metadata.create_all(self.engine)
        now=datetime.combine(business_today(),time(10))
        with SessionLocal() as db:
            db.add_all([User(id=1,full_name='Manager',email='manager@example.test',password_hash=self.password_hash,role='MANAGEMENT'),
                        User(id=2,full_name='Sales A',email='a@example.test',password_hash=self.password_hash,role='SALES'),
                        User(id=3,full_name='Sales B',email='b@example.test',password_hash=self.password_hash,role='SALES')])
            db.add_all([Client(id=1,business_name='<script>alert(1)</script>',phone='111',area='Gota'),
                        Client(id=2,business_name='=HYPERLINK("https://example.test")',phone='222',area='Rajkot'),
                        Client(id=3,business_name='Recent Salon',phone='333')]);db.flush()
            for aid,cid,days in [(1,1,8),(2,1,7),(3,2,9),(4,3,3)]:
                db.add(Activity(id=aid,client_id=cid,activity_type='FIELD_VISIT',status='COMPLETED',completed_at=now-timedelta(days=days),created_at=now-timedelta(days=days),assigned_user_id=2))
                db.add(Visit(activity_id=aid,visit_type='SALES',check_in_at=now-timedelta(days=days),check_out_at=now-timedelta(days=days)+timedelta(minutes=20)))
            db.add(Activity(id=5,client_id=1,activity_type='COLLECTION',status='OPEN',created_at=now-timedelta(days=2)))
            db.add(CollectionActivity(id=1,activity_id=5,channel='TELECALLING'));db.flush()
            for pid,days in [(1,0),(2,-1),(3,1)]:
                db.add(PromiseToPay(id=pid,client_id=1,collection_activity_id=1,promised_amount=5000,promise_date=now+timedelta(days=days),status='OPEN'))
            db.add(Activity(id=6,client_id=1,activity_type='FIELD_VISIT',status='CREATED',assigned_user_id=2))
            db.add(Visit(activity_id=6,visit_type='SALES'));db.commit()
        self.client=TestClient(app)

    def tearDown(self):
        self.client.close();Base.metadata.drop_all(self.engine)

    def login(self,who='manager'):
        email={'manager':'manager@example.test','a':'a@example.test','b':'b@example.test'}[who]
        response=self.client.post('/login',data={'email':email,'password':'Test-only-password-123'},follow_redirects=False)
        self.assertEqual(response.status_code,303)

    def workbook(self,response):
        self.assertEqual(response.status_code,200)
        return list(load_workbook(BytesIO(response.content)).active.values)

    def test_courtesy_latest_visit_once_per_client(self):
        self.login();rows=self.workbook(self.client.get('/reports/export?kind=courtesy'))
        self.assertEqual(len(rows),3)
        self.assertEqual({r[0] for r in rows[1:]},{1,2})
        salon1=next(r for r in rows[1:] if r[0]==1)
        self.assertEqual(salon1[6],0)

    def test_courtesy_filter_and_export_agree(self):
        self.login();url='/operations?kind=courtesy&q=Rajkot'
        self.assertIn('1 matching records',self.client.get(url).text)
        rows=self.workbook(self.client.get('/reports/export?kind=courtesy&q=Rajkot'))
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][0],2)
        response=self.client.get('/reports/export?kind=courtesy&start='+str(business_today())+'&end='+str(business_today()))
        self.assertEqual(len(self.workbook(response)),2)

    def test_ptp_due_today_and_overdue_not_future(self):
        self.login();rows=self.workbook(self.client.get('/reports/export?kind=ptp'))
        self.assertEqual({r[0] for r in rows[1:]},{1,2})
        self.assertEqual({r[6] for r in rows[1:]},{'TODAY','OVERDUE'})
        today=str(business_today());rows=self.workbook(self.client.get(f'/reports/export?kind=ptp&start={today}&end={today}'))
        self.assertEqual(len(rows),2);self.assertEqual(rows[1][0],1)

    def test_html_and_excel_text_are_not_executed(self):
        self.login();html=self.client.get('/operations?kind=courtesy').text
        self.assertNotIn('<script>alert(1)</script>',html);self.assertIn('&lt;script&gt;',html)
        response=self.client.get('/reports/export?kind=courtesy&q=222')
        book=load_workbook(BytesIO(response.content));self.assertEqual(book.active['B2'].data_type,'s')
        self.assertEqual(book.active['B2'].value,'=HYPERLINK("https://example.test")')
        self.assertNotIn('<script>alert(1)</script>',self.client.get('/clients/1/timeline').text)

    def test_reports_and_timeline_role_permissions(self):
        for path in ['/operations','/reports/export','/clients/1/timeline']:
            self.assertEqual(self.client.get(path,follow_redirects=False).status_code,303)
        self.login('a')
        for path in ['/operations','/reports/export','/clients/1/timeline']:
            self.assertEqual(self.client.get(path).status_code,403)

    def test_invalid_report_dates(self):
        self.login()
        self.assertEqual(self.client.get('/reports/export?kind=nope').status_code,400)
        self.assertEqual(self.client.get('/operations?start=2026-10-09&end=2026-10-08').status_code,400)
        self.assertEqual(self.client.get('/operations?start=invalid').status_code,422)
        self.assertEqual(self.client.get('/clients/999/timeline').status_code,404)

    def test_other_employee_cannot_view_or_mutate_visit(self):
        self.login('b')
        self.assertEqual(self.client.get('/visits/6').status_code,403)
        for suffix,data in [('start',{'latitude':23,'longitude':72}),('request-end-otp',{'notes':'x'}),('verify-end-otp',{'otp':'123456'})]:
            self.assertEqual(self.client.post('/visits/6/'+suffix,data=data).status_code,403)
        with SessionLocal() as db:
            self.assertEqual(db.get(Activity,6).status,'CREATED')
            self.assertIsNone(db.query(Visit).filter_by(activity_id=6).one().check_in_at)

    def test_valid_gps_checkin_and_immutable_timestamp(self):
        self.login('a');data={'latitude':23,'longitude':72}
        self.assertEqual(self.client.post('/visits/6/start',data=data).status_code,200)
        with SessionLocal() as db:
            original=db.query(Visit).filter_by(activity_id=6).one().check_in_at
            self.assertEqual(db.get(Client,1).latitude,23)
        self.assertEqual(self.client.post('/visits/6/start',data=data).status_code,409)
        with SessionLocal() as db:
            self.assertEqual(db.query(Visit).filter_by(activity_id=6).one().check_in_at,original)
        self.assertEqual(self.client.get('/visits/6').status_code,200)
        self.assertEqual(self.client.get('/visits').status_code,200)

    def test_bad_gps_and_missing_visits_do_not_write(self):
        self.login('a')
        for lat,lon in [(100,72),(23,200),('nan',72),('inf',72)]:
            self.assertEqual(self.client.post('/visits/6/start',data={'latitude':lat,'longitude':lon}).status_code,400)
        for suffix,data in [('start',{'latitude':23,'longitude':72}),('request-end-otp',{'notes':'x'}),('verify-end-otp',{'otp':'123456'})]:
            self.assertEqual(self.client.post('/visits/999/'+suffix,data=data).status_code,404)
        self.assertEqual(self.client.post('/visits/create',data={'client_id':999,'visit_type':'SALES'}).status_code,404)
        self.assertEqual(self.client.post('/visits/create',data={'client_id':1,'visit_type':'BAD'}).status_code,400)
        with SessionLocal() as db:self.assertEqual(db.get(Activity,6).status,'CREATED')

    def test_gps_radius_enforced(self):
        self.login('a')
        with SessionLocal() as db:
            c=db.get(Client,1);c.latitude=23;c.longitude=72;db.commit()
        self.assertEqual(self.client.post('/visits/6/start',data={'latitude':23.003,'longitude':72}).status_code,409)
        self.assertEqual(self.client.post('/visits/6/start',data={'latitude':23.001,'longitude':72}).status_code,200)

    def test_production_does_not_expose_or_queue_test_otp(self):
        self.login('a');self.client.post('/visits/6/start',data={'latitude':23,'longitude':72})
        with patch.dict(os.environ,{'ESONE_ENV':'production'}):
            response=self.client.post('/visits/6/request-end-otp',data={'notes':'Visit done'})
        self.assertEqual(response.status_code,503);self.assertNotIn('TEST OTP',response.text)
        with SessionLocal() as db:self.assertEqual(db.query(VerificationOTP).count(),0)

    def test_test_only_otp_completion_and_replay_rejected(self):
        import re
        self.login('a');self.client.post('/visits/6/start',data={'latitude':23,'longitude':72})
        with patch.dict(os.environ,{'ESONE_ENV':'test'}):
            response=self.client.post('/visits/6/request-end-otp',data={'notes':'Visit done'})
        code=re.search(r'TEST OTP: (\d+)',response.text).group(1)
        self.assertEqual(self.client.post('/visits/6/verify-end-otp',data={'otp':code},follow_redirects=False).status_code,303)
        self.assertEqual(self.client.post('/visits/6/verify-end-otp',data={'otp':code},follow_redirects=False).status_code,409)
        with SessionLocal() as db:
            self.assertEqual(db.get(Activity,6).status,'COMPLETED')
            self.assertEqual(db.get(Activity,6).notes,'Visit done')
            self.assertEqual(db.query(AuditLog).count(),2)

    def test_readonly_pages_and_exports_leave_records_unchanged(self):
        self.login()
        with SessionLocal() as db:
            before={t.name:db.execute(t.select()).all() for t in Base.metadata.sorted_tables}
        for path in ['/operations','/operations?kind=ptp','/reports/export?kind=courtesy','/reports/export?kind=visits','/clients/1/timeline','/visits']:
            self.assertEqual(self.client.get(path).status_code,200)
        with SessionLocal() as db:
            after={t.name:db.execute(t.select()).all() for t in Base.metadata.sorted_tables}
        self.assertEqual(before,after)

    def test_readiness_is_read_only_and_missing_schema_fails_clearly(self):
        with patch('database.db.engine',self.engine):
            self.assertEqual(self.client.get('/health/ready').json(),{'status':'ready','service':'ESONE'})
        blank=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
        try:
            with patch('database.db.engine',blank):
                self.assertEqual(self.client.get('/health/ready').status_code,503)
        finally:blank.dispose()

    def test_database_failure_does_not_leak_details_or_initialize_schema(self):
        from sqlalchemy.exc import OperationalError
        with patch('main.SessionLocal',side_effect=OperationalError('secret SQL',{},Exception('secret-db-url'))):
            response=self.client.post('/login',data={'email':'x@example.test','password':'test'})
        self.assertEqual(response.status_code,503)
        self.assertNotIn('secret',response.text)

    def test_mobile_has_actual_visit_count_and_no_fake_revenue(self):
        self.login('a');self.client.post('/visits/6/start',data={'latitude':23,'longitude':72})
        response=self.client.get('/mobile')
        self.assertEqual(response.status_code,200)
        self.assertIn('<div class="number">1</div>',response.text)
        self.assertNotIn('₹0',response.text)

    def test_old_public_session_secret_cannot_forge_access(self):
        import base64,json
        from itsdangerous import TimestampSigner
        payload=base64.b64encode(json.dumps({'user_id':1,'role':'SUPER_ADMIN'}).encode())
        token=TimestampSigner('ES1-DEVELOPMENT-SECRET-CHANGE-LATER').sign(payload).decode()
        self.client.cookies.set('session',token)
        self.assertEqual(self.client.get('/operations',follow_redirects=False).status_code,303)

    def test_long_password_and_bad_hash_fail_without_exception(self):
        from modules.security import verify_password
        self.assertFalse(verify_password('x'*73,self.password_hash))
        self.assertFalse(verify_password('any','invalid hash'))

    def test_search_attribute_is_escaped(self):
        self.login()
        response=self.client.get('/clients',params={'q':'" autofocus onfocus="alert(1)'})
        self.assertEqual(response.status_code,200)
        self.assertNotIn('value="" autofocus',response.text)
