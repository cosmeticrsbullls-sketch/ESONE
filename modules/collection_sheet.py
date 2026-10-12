"""Read outstanding-bill sheets without inventing dated payment receipts."""
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

MAX_ROWS=5000

def parse_collection_sheet(data):
    # Excel formatting can extend to one million rows; never trust dimensions.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist())>200*1024*1024:raise ValueError('Workbook is too large after decompression.')
    workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    records=[];warnings=[];ignored=0
    try:
        for sheet in workbook:
            header=None
            for row_number,row in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or MAX_ROWS,MAX_ROWS),max_col=25,values_only=True),1):
                if header is None:
                    labels=[re.sub(r'\s+',' ',str(x or '').strip().upper()) for x in row]
                    if 'PARTY NAME' in labels and 'BALANCE' in labels and 'BILL AMOUNT' in labels:
                        header={value:index for index,value in enumerate(labels) if value}
                    continue
                def cell(*labels):
                    for label in labels:
                        if label in header:return row[header[label]]
                name=cell('PARTY NAME');amount=cell('BILL AMOUNT');balance=cell('BALANCE')
                if not name:
                    if balance is not None:ignored+=1
                    continue
                if amount is None:
                    warnings.append(f'{sheet.title}, row {row_number}: missing bill amount; row was not imported.')
                    continue
                try:
                    total=Decimal(str(amount).replace(',',''));outstanding=Decimal(str(balance).replace(',','')) if balance is not None else None
                    if not total.is_finite() or total<0 or outstanding is None or not outstanding.is_finite() or outstanding<0 or outstanding>total:raise ValueError()
                    if total!=total.quantize(Decimal('.01')) or outstanding!=outstanding.quantize(Decimal('.01')):raise ValueError()
                except (ValueError,InvalidOperation):
                    warnings.append(f'{sheet.title}, row {row_number}: check bill amount and balance.');continue
                raw_date=cell('DATE')
                bill_date=raw_date.date().isoformat() if isinstance(raw_date,datetime) else raw_date.isoformat() if isinstance(raw_date,date) else ''
                raw_reference=cell('BILL NO.','BILL NO')
                reference=str(int(raw_reference)) if isinstance(raw_reference,(int,float)) and raw_reference==int(raw_reference) else str(raw_reference or '').strip()
                raw_phone=cell('CONTACET NO','CONTACT NO','CONTACT NO.')
                phone=str(int(raw_phone)) if isinstance(raw_phone,(int,float)) else str(raw_phone or '').strip()
                records.append(dict(sheet=sheet.title,row=row_number,customer_name=str(name).strip(),area=str(cell('AREA') or '').strip(),phone=phone,reference=reference,document_date=bill_date,bill_amount=str(total),outstanding=str(outstanding),remarks=str(cell('REMARKS') or '').strip()))
            if (sheet.max_row or 0)>MAX_ROWS:
                warnings.append(f'{sheet.title}: only the first {MAX_ROWS} rows were scanned; confirm there are no later business records. Blank/formatted rows are not imported.')
        if not records:raise ValueError('No supported outstanding bill rows found.')
        return dict(format='OUTSTANDING_SHEET',records=records,ignored_summary_rows=ignored,warnings=warnings,total_outstanding=str(sum(Decimal(r['outstanding']) for r in records)))
    finally:workbook.close()
